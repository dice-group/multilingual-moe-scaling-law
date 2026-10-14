"""Presentation helpers for the publication-grade Stage 1 A--D release.

The functions in this module are deliberately downstream-only: callers supply
release metadata and already-built pandas tables.  No raw result artifact is
discovered or interpreted here.  Parquet remains the canonical analytical
format; the Excel workbook is a documented convenience view.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from copy import deepcopy
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
from typing import Any


EXCEL_MAX_DATA_ROWS = 1_048_575
EXCEL_MAX_COLUMNS = 16_384
EXCEL_MAX_CELL_CHARACTERS = 32_767

PROVENANCE_CLASSES: tuple[dict[str, str], ...] = (
    {
        "provenance_class": "observed",
        "definition": (
            "Read directly from an immutable run, evaluation, scheduler, or "
            "telemetry artifact."
        ),
    },
    {
        "provenance_class": "configured",
        "definition": (
            "Declared by a frozen study specification, resolved recipe, or "
            "launch manifest."
        ),
    },
    {
        "provenance_class": "planned",
        "definition": (
            "A pre-run target, budget, or forecast; never evidence that work "
            "was executed."
        ),
    },
    {
        "provenance_class": "derived",
        "definition": (
            "Recomputed from observed/configured values by release code; the "
            "source columns and method must remain traceable."
        ),
    },
)


def _catalog_entry(
    table_name: str,
    category: str,
    grain: str,
    purpose: str,
    primary_key: Sequence[str],
    key_columns: Sequence[str],
    caveats: str,
    *,
    excel_default: bool,
) -> dict[str, Any]:
    return {
        "table_name": table_name,
        "category": category,
        "grain": grain,
        "purpose": purpose,
        "primary_key": list(primary_key),
        "key_columns": list(key_columns),
        "caveats": caveats,
        "excel_default": excel_default,
    }


_TABLE_CATALOG: tuple[dict[str, Any], ...] = (
    _catalog_entry(
        "release_summary",
        "release",
        "one row per reporting scope (A, B, C, D, or ALL)",
        "Compact per-ladder and all-release counts, compute, and energy coverage.",
        ("scope",),
        (
            "scope",
            "independent_trajectories",
            "checkpoint_readouts",
            "production_training_jobs",
            "allocated_gpu_hours_including_retries",
        ),
        "The manifest and its hashes, not this convenience row, are authoritative.",
        excel_default=True,
    ),
    _catalog_entry(
        "studies",
        "design",
        "one row per frozen ladder protocol",
        "Protocol/specification identity and fixed training-policy fields.",
        ("protocol_id", "ladder"),
        (
            "protocol_id",
            "study_id",
            "ladder",
            "spec_path",
            "spec_sha256",
            "model_seed",
            "dataset_shuffle_seed",
            "provenance_class",
        ),
        "Configured values must not be presented as runtime observations.",
        excel_default=True,
    ),
    _catalog_entry(
        "source_roots",
        "lineage",
        "one row per primary or recovery run root",
        "Enumerates the source boundary and role of each admitted run root.",
        ("source_root_id",),
        ("source_root_id", "portable_path", "role", "exists"),
        "Recovery/offload roots are evidence sources, not extra scientific replicates.",
        excel_default=True,
    ),
    _catalog_entry(
        "model_configs",
        "design",
        "one row per ladder-by-expert-count model configuration",
        "Resolved dense/MoE architecture, parameter, routing, and recipe fields.",
        ("model_config_id",),
        (
            "model_config_id",
            "ladder",
            "experts",
            "architecture_type",
            "num_layers",
            "hidden_size",
            "sequence_length",
        ),
        "Declared active and runtime total parameter counts are different quantities.",
        excel_default=True,
    ),
    _catalog_entry(
        "trajectories",
        "design",
        "one row per ladder-by-expert-by-language trajectory",
        "The independent configured training trajectory and its seed identity.",
        ("trajectory_id",),
        (
            "trajectory_id",
            "ladder",
            "experts",
            "language_count",
            "model_seed",
            "dataset_shuffle_seed",
            "independent_seed_replicates",
        ),
        (
            "Checkpoint and language rows within a trajectory are repeated "
            "measurements, not independent seed replicates."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "readouts",
        "validation",
        "one row per admitted schedule-aware checkpoint readout",
        "Central A--D outcome index joining trajectory, checkpoint, and evaluation.",
        ("readout_id",),
        (
            "readout_id",
            "trajectory_id",
            "ladder",
            "experts",
            "language_count",
            "actual_tokens",
            "primary_metric_value",
            "evaluation_manifest_path",
        ),
        (
            "Use actual optimizer-step-aligned tokens. Nominal horizons are "
            "targets and can differ slightly."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "validation_aggregates",
        "validation",
        "one row per readout-by-aggregate definition",
        "Macro/micro and trained/zero-shot validation summaries.",
        ("readout_id", "aggregate_kind", "cohort"),
        (
            "readout_id",
            "aggregate_kind",
            "cohort",
            "metric_name",
            "metric_value",
            "aggregate_language_count",
            "aggregate_token_count",
        ),
        "The aggregation rule and denominator must accompany every comparison.",
        excel_default=True,
    ),
    _catalog_entry(
        "validation_language",
        "validation",
        "one row per readout-by-validation-language",
        "Raw per-language NLL and language role for hierarchical analysis.",
        ("readout_id", "language"),
        (
            "readout_id",
            "language",
            "metric_role",
            "atlas_tao_unigram_adjusted_loss",
            "raw_nll",
            "token_count",
            "validation_input_identity_sha256",
            "input_sha256",
            "input_sha256_semantics",
            "language_artifact_path",
        ),
        (
            "Languages in the same readout and checkpoints in the same "
            "trajectory are statistically dependent."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "validation_scope",
        "validation",
        "one row per ladder-by-language-tier validation coverage cell",
        (
            "Explicit trained and zero-shot language coverage matrix for every "
            "A-D ladder and K tier."
        ),
        ("ladder", "language_count"),
        (
            "validation_scope_id",
            "ladder",
            "language_count",
            "readout_count",
            "trained_languages_per_readout",
            "zero_shot_languages_per_readout",
            "unique_validation_languages",
            "scope_status",
        ),
        (
            "Only K1 readouts contain zero-shot language outcomes. K12, K50, "
            "and K200 evaluate their trained tier only, so cross-tier zero-shot "
            "comparisons are unavailable."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "validation_candidates",
        "validity",
        "one row per candidate evaluation artifact",
        "Candidate selection, duplicate handling, and validity verdicts.",
        ("readout_id", "candidate_manifest_path"),
        (
            "readout_id",
            "candidate_manifest_path",
            "ledger_valid",
            "selected",
            "semantically_matches_selected",
            "selection_reason",
        ),
        (
            "Primary, recovery, and offload copies are provenance alternatives, "
            "not independent results."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "validation_artifacts",
        "lineage",
        "one row per validation artifact file",
        "File-level validation inventory with hashes and selected status.",
        ("artifact_id",),
        ("artifact_id", "readout_id", "artifact_type", "path", "sha256"),
        "High-cardinality lineage table; inspect in Parquet for complete coverage.",
        excel_default=False,
    ),
    _catalog_entry(
        "validation_provenance",
        "lineage",
        "one row per validation subject-to-object provenance edge",
        "Machine-readable lineage graph for readouts and validation artifacts.",
        ("provenance_id",),
        (
            "provenance_id",
            "readout_id",
            "relation",
            "subject_path",
            "object_path",
            "source_path",
            "source_sha256",
        ),
        "Edges express lineage, not causal or statistical independence.",
        excel_default=False,
    ),
    _catalog_entry(
        "training_jobs",
        "training",
        "one row per logical production training job",
        "Training topology, trunk/branch role, targets, outcomes, and manifests.",
        ("training_job_id",),
        (
            "training_job_id",
            "trajectory_id",
            "job_role",
            "parent_training_job_id",
            "target_tokens",
            "actual_end_tokens",
            "status",
        ),
        (
            "Shared trunks must be counted once in physical compute totals, "
            "even when several readouts inherit their checkpoint."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "training_segments",
        "training",
        "one row per logical job segment/log",
        "Segment boundaries, batch settings, source logs, and hashes.",
        ("segment_id",),
        (
            "segment_id",
            "training_job_id",
            "segment_index",
            "segment_kind",
            "actual_start_tokens",
            "actual_end_tokens",
            "log_sha256",
        ),
        "A segment is not an independent run; order by job and segment index.",
        excel_default=True,
    ),
    _catalog_entry(
        "training_metrics",
        "training",
        "one row per parsed Megatron training-log observation",
        "Fine-grained loss, throughput, timing, parameter, and MFU observations.",
        (
            "training_job_id",
            "segment_id",
            "event_type",
            "source_line_number",
        ),
        (
            "training_job_id",
            "segment_id",
            "event_type",
            "source_line_number",
            "timestamp",
            "iteration",
            "consumed_tokens",
            "lm_loss",
            "elapsed_time_ms",
            "tokens_per_second",
            "model_tflops_per_gpu",
        ),
        (
            "Serial correlation is strong. Hardware throughput/MFU is diagnostic, "
            "not an independent model-quality outcome."
        ),
        excel_default=False,
    ),
    _catalog_entry(
        "training_summary",
        "training",
        "one row per logical production training job",
        "Robust training and telemetry summaries with completeness indicators.",
        ("training_job_id",),
        (
            "training_job_id",
            "trajectory_id",
            "parsed_training_metric_rows",
            "training_metric_coverage_fraction",
            "mean_step_ms",
            "median_step_ms",
            "mean_tokens_per_second",
            "mean_model_tflops_per_gpu",
            "mean_mfu_pct_of_configured_peak",
        ),
        "Aggregation windows and telemetry coverage must be reported with summaries.",
        excel_default=True,
    ),
    _catalog_entry(
        "gpu_samples",
        "compute",
        "one row per timestamp-by-GPU telemetry sample",
        "GPU utilization, memory, temperature, power, and board-energy inputs.",
        ("training_job_id", "timestamp", "gpu_uuid"),
        (
            "training_job_id",
            "timestamp",
            "gpu_uuid",
            "gpu_utilization_pct",
            "memory_used_mib",
            "board_power_draw_w",
            "sample_valid",
            "seconds_since_previous_sample",
        ),
        (
            "GPU-board power is not total facility/system energy and excludes "
            "CPU, network, storage, cooling, and PUE."
        ),
        excel_default=False,
    ),
    _catalog_entry(
        "training_artifacts",
        "lineage",
        "one row per training log, telemetry file, or training manifest",
        "Hash inventory for the raw inputs admitted to training tables.",
        ("artifact_type", "object_id", "path"),
        ("artifact_type", "object_id", "path", "sha256", "size_bytes"),
        "Large immutable inputs remain outside the release and are referenced by hash.",
        excel_default=False,
    ),
    _catalog_entry(
        "data_exposure",
        "data",
        "one row per readout-by-trained-language",
        "Exact realized token exposure and reuse relative to source inventory.",
        ("readout_id", "language"),
        (
            "readout_id",
            "language",
            "metric_role",
            "realized_self_exposure_tokens",
            "exact_unique_language_tokens_including_eod",
            "effective_exposure_passes",
        ),
        (
            "Exact blend replay is authoritative; nominal D/K is not an exact "
            "per-language exposure."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "routing_layers",
        "routing",
        "one row per readout-by-language-by-MoE-layer routing record",
        "Expert token allocation and routing imbalance diagnostics.",
        ("readout_id", "language", "layer_index"),
        (
            "readout_id",
            "language",
            "layer_index",
            "expert_token_counts_json",
        ),
        (
            "Routing observations are conditional on checkpoint, language, "
            "batch, and layer; they are not independent replicates."
        ),
        excel_default=False,
    ),
    _catalog_entry(
        "job_attempts",
        "compute",
        "one row per discovered logical Slurm submission attempt",
        "Attempt lineage joining run artifacts to scheduler accounting records.",
        ("attempt_uid",),
        (
            "attempt_uid",
            "logical_job_id",
            "slurm_job_id",
            "attempt_index",
            "accounting_status",
        ),
        "Retries and arrays must not be conflated with independent scientific runs.",
        excel_default=True,
    ),
    _catalog_entry(
        "slurm_allocations",
        "compute",
        "one row per normalized Slurm job/array-task/step accounting record",
        "Observed allocation, elapsed time, state, node, GPU, and TRES accounting.",
        ("allocation_uid",),
        (
            "allocation_uid",
            "attempt_uid",
            "logical_job_id",
            "slurm_record_id",
            "slurm_job_id",
            "state",
            "elapsed_seconds",
            "allocated_gpus",
            "allocated_gpu_type_counts",
        ),
        (
            "Scheduler records describe allocations. Do not infer application "
            "utilization from allocated GPU-hours alone."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "compute_summary",
        "compute",
        "one row per declared compute aggregation unit",
        "Deduplicated GPU-hours, wall time, hardware, utilization, and energy boundary.",
        ("training_job_id",),
        (
            "training_job_id",
            "trajectory_id",
            "attempt_count",
            "accounted_attempt_count",
            "allocated_gpu_hours_including_retries",
            "allocated_gpu_type_counts_across_attempts_json",
            "board_energy_estimate_kwh",
        ),
        (
            "Physical-job totals must deduplicate shared trunks and retries; "
            "board energy and total-system energy are distinct."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "validity_events",
        "validity",
        "one row per validity decision or exclusion event",
        "Predeclared technical-validity decisions and their evidence.",
        ("object_type", "object_id", "issue_type"),
        (
            "object_type",
            "object_id",
            "issue_type",
            "severity",
            "verdict",
            "evidence_path",
        ),
        "Scientific outcomes must not be excluded merely because their loss is poor.",
        excel_default=True,
    ),
    _catalog_entry(
        "qa_issues",
        "validity",
        "one row per release-quality issue",
        "Failed or warning-level schema, lineage, count, and accounting checks.",
        ("qa_issue_id",),
        (
            "qa_issue_id",
            "check_id",
            "severity",
            "status",
            "details",
            "observed",
            "expected",
        ),
        "A release with unresolved error-level issues must not be marked sealed.",
        excel_default=True,
    ),
    _catalog_entry(
        "analysis_registry",
        "analysis",
        "one row per frozen or exploratory analysis artifact",
        "Registers notebooks, fit receipts, simulations, and protocol decisions.",
        ("analysis_artifact_id",),
        (
            "analysis_artifact_id",
            "artifact_role",
            "status",
            "scope_disclosure",
            "path",
            "sha256",
            "sequence_order",
            "depends_on",
        ),
        (
            "Frozen confirmatory analyses and post-hoc exploratory analyses "
            "must remain visibly distinct."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "neurips_checklist",
        "release",
        "one row per checklist topic",
        "Machine-readable evidence and remaining manuscript actions.",
        ("checklist_id",),
        (
            "checklist_id",
            "topic",
            "release_support",
            "paper_answer",
            "evidence",
            "action_required",
        ),
        "This supports, but cannot replace, the authors' paper-level checklist answers.",
        excel_default=False,
    ),
)


_TABLE_CATALOG += (
    _catalog_entry(
        "training_configs",
        "training",
        "one row per logical production training job",
        "Full resolved recipe, batch, schedule, data, and artifact configuration snapshot.",
        ("training_job_id",),
        (
            "training_job_id",
            "trajectory_id",
            "job_role",
            "recipe",
            "backend",
            "sequence_length",
            "global_batch_size",
            "training_manifest_path",
            "training_manifest_sha256",
        ),
        (
            "Configured fields and observed manifest identity have distinct "
            "field-level provenance; retain field_provenance_json."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "runtime_parameters",
        "training",
        "one row per unique job-segment parameter-scope/rank observation",
        "Deduplicated runtime parameter-count observations parsed from training logs.",
        (
            "training_job_id",
            "segment_id",
            "parameter_scope",
            "source_line_number",
        ),
        (
            "training_job_id",
            "trajectory_id",
            "segment_id",
            "parameter_scope",
            "source_line_number",
            "parameter_count",
            "tensor_model_parallel_rank",
            "pipeline_model_parallel_rank",
            "source_path",
        ),
        "Runtime total, active, and non-embedding parameter scopes are not interchangeable.",
        excel_default=True,
    ),
    _catalog_entry(
        "production_artifact_verifications",
        "lineage",
        "one row per production training job, artifact type, and path",
        (
            "Independent SHA-256 verification of every execution-bound "
            "training manifest and resolved configuration."
        ),
        ("training_job_id", "artifact_type", "path"),
        (
            "training_job_id",
            "artifact_type",
            "path",
            "expected_sha256",
            "recomputed_sha256",
            "status",
        ),
        (
            "Verification establishes current byte identity against the "
            "execution receipt; it does not turn a configured value into an "
            "observed runtime measurement."
        ),
        excel_default=False,
    ),
    _catalog_entry(
        "readout_training_bindings",
        "lineage",
        "one row per admitted schedule-aware checkpoint readout",
        (
            "Fail-closed binding from each validation readout to its production "
            "training job and checkpoint iteration."
        ),
        ("readout_id",),
        (
            "readout_id",
            "training_job_id",
            "checkpoint_iteration_exists_now",
            "status",
            "errors",
        ),
        (
            "This table verifies lineage and current checkpoint presence; it "
            "does not create an additional outcome or replicate."
        ),
        excel_default=False,
    ),
    _catalog_entry(
        "dataset_release",
        "data",
        "one row per frozen corpus release",
        "Corpus release contract, readiness, manifest identity, and selection inputs.",
        ("dataset_release_id",),
        (
            "dataset_release_id",
            "schema_version",
            "release_status",
            "ready",
            "language_count",
            "manifest_path",
            "manifest_file_sha256",
        ),
        "The manifest and per-language inventory, not this summary alone, define the corpus.",
        excel_default=True,
    ),
    _catalog_entry(
        "corpus_languages",
        "data",
        "one row per corpus-release-by-language inventory record",
        "Per-language train/validation selection, uniqueness, coverage, paths, and hashes.",
        ("dataset_release_id", "language"),
        (
            "dataset_release_id",
            "language",
            "exact_unique_language_tokens_including_eod",
            "train_selected_tokens",
            "validation_selected_tokens",
            "train_sha256",
            "validation_sha256",
            "corpus_manifest_sha256",
        ),
        "Unique-token and selected-token counts answer different exposure questions.",
        excel_default=True,
    ),
    _catalog_entry(
        "launch_provenance",
        "lineage",
        "one row per production execution family (AB, C, or D)",
        (
            "Execution-bound launch manifest, code revision, runtime profile, "
            "and sealed-corpus identity."
        ),
        ("execution_id",),
        (
            "execution_id",
            "source_family",
            "launch_manifest_path",
            "launch_manifest_file_sha256",
            "orchestration_commit",
            "bridge_commit",
            "data_seal_sha256",
        ),
        (
            "This records what was executed. A later checkout or planning "
            "configuration must not replace the execution-bound evidence."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "execution_inputs",
        "lineage",
        "one row per execution-by-declared fixed input",
        ("Launch-declared input identity and current preservation/hash verification."),
        ("execution_id", "declared_path"),
        (
            "execution_id",
            "source_family",
            "declared_path",
            "resolved_path",
            "expected_sha256",
            "current_sha256",
            "current_hash_matches_execution",
        ),
        (
            "The expected execution hash is authoritative. Current equality is "
            "a preservation check, not permission to substitute a later file."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "hardware_contract",
        "compute",
        "one row per frozen hardware-efficiency denominator contract",
        (
            "Hardware identity, configured peak denominator, MFU formula, and "
            "source seal."
        ),
        ("hardware_contract_id",),
        (
            "hardware_contract_id",
            "accelerator_family",
            "precision",
            "peak_tflops_per_gpu",
            "denominator_kind",
            "source_document_sha256",
            "mfu_formula",
        ),
        (
            "Configured theoretical peak is a denominator, not achieved "
            "throughput or end-to-end facility efficiency."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "languages",
        "data",
        "one row per canonical production language",
        (
            "Production language/script identity, corpus inventory, and "
            "explicitly sourced family/resource metadata."
        ),
        ("language",),
        (
            "language",
            "iso639_3",
            "production_script",
            "k200_rank",
            "family",
            "resource_availability",
            "legacy_metadata_join_method",
            "stale_script_mismatch",
        ),
        (
            "Production script comes from the sealed v2 tier. Legacy metadata "
            "is joined only by an unambiguous ISO-639-3 prefix and its legacy "
            "script is never promoted to production identity."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "language_tier_membership",
        "data",
        "one row per ordered language-tier membership",
        ("Canonical K1/K12/K50/K200 membership, rank, source, and membership seal."),
        ("language_count", "language"),
        (
            "language_set_id",
            "tier",
            "language_count",
            "tier_rank",
            "language",
            "ordered_membership_sha256",
            "source_file_sha256",
        ),
        (
            "Tier rows encode design membership and order; they are not "
            "additional language observations."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "language_metadata_sources",
        "lineage",
        "one row per language-dimension source artifact",
        "Hash-sealed registry of production tier and legacy metadata sources.",
        ("language_metadata_source_id",),
        (
            "language_metadata_source_id",
            "source_type",
            "path",
            "file_sha256",
            "embedded_sha256",
            "language_count",
        ),
        (
            "Legacy family/name/resource metadata and current production tier "
            "identity have different evidentiary roles."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "validation_cohorts",
        "validation",
        "one row per distinct aggregate-kind/cohort/membership definition",
        (
            "Deduplicated validation cohort membership, membership hash, and "
            "readout coverage."
        ),
        ("validation_cohort_id",),
        (
            "validation_cohort_id",
            "aggregate_kind",
            "cohort",
            "language_count",
            "membership_sha256",
            "readout_count",
        ),
        (
            "This is a cohort definition table, not an outcome table. Join "
            "outcomes through validation_aggregates."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "readout_compute_attribution",
        "compute",
        "one row per schedule-aware checkpoint readout",
        (
            "Logical readout cost attribution combining its parent-trunk prefix "
            "and branch work."
        ),
        ("readout_id",),
        (
            "readout_id",
            "trajectory_id",
            "branch_training_job_id",
            "parent_trunk_training_job_id",
            "parent_trunk_prefix_tokens",
            "logical_readout_attributed_productive_gpu_hours",
            "physical_aggregation_allowed",
        ),
        (
            "This is attributed logical cost, not additive physical study "
            "compute. Shared trunk prefixes recur across readouts; never sum "
            "this table to estimate physical compute."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "unique_physical_job_compute",
        "compute",
        "one row per unique physical production training job",
        ("Deduplicated job-level productive, allocated, and board-energy accounting."),
        ("physical_compute_unit_id",),
        (
            "physical_compute_unit_id",
            "training_job_id",
            "trajectory_id",
            "job_role",
            "estimated_productive_gpu_hours",
            "allocated_gpu_hours_including_retries",
            "physical_deduplication_key",
        ),
        (
            "Each physical training job appears once. Allocated GPU-hours can "
            "include retries and are not application utilization."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "unique_physical_study_compute",
        "compute",
        "one row per deduplicated Stage 1 A-D physical training scope",
        (
            "Study-level physical training totals, coverage, and explicit "
            "non-additivity guard for logical attribution."
        ),
        ("physical_compute_scope_id",),
        (
            "physical_compute_scope_id",
            "unique_physical_training_job_count",
            "unique_physical_trunk_count",
            "unique_physical_branch_count",
            "estimated_productive_gpu_hours_unique_jobs",
            "observed_allocated_gpu_hours_including_retries_unique_jobs",
            "logical_attribution_sum_status",
        ),
        (
            "This is the additive training-compute total. It excludes logical "
            "readout-attribution sums and discloses metric coverage."
        ),
        excel_default=True,
    ),
    _catalog_entry(
        "physical_compute_categories",
        "compute",
        "one row per mutually exclusive workload category plus all-records total",
        (
            "Deduplicated scheduler compute split by core training, selected or "
            "recovery evaluation, data, control, systems, and other work."
        ),
        ("workload_category",),
        (
            "workload_category",
            "discovered_attempts",
            "accounted_attempts",
            "missing_accounting_attempts",
            "unique_physical_allocation_records",
            "allocated_gpu_hours_all_states",
            "allocated_cpu_hours_all_states",
        ),
        (
            "Categories describe workload purpose, not scientific independence. "
            "The all-records row is a total and must not be summed with its "
            "component categories."
        ),
        excel_default=True,
    ),
)

_DEFAULT_DISCLOSURES: tuple[dict[str, str], ...] = (
    {
        "disclosure_id": "single_seed",
        "category": "statistical scope",
        "severity": "limitation",
        "title": "One configured seed per design cell/trajectory",
        "statement": (
            "The independent unit is one full training trajectory: 64 "
            "trajectories in the A-D design, with one configured model/data "
            "seed per cell. Checkpoint readouts and per-language rows are "
            "repeated measures, not independent seed replicates."
        ),
        "analysis_consequence": (
            "Cluster inference at trajectory_id. Avoid seed-general claims; "
            "use descriptive, paired, or hierarchical analyses that retain "
            "trajectory identity."
        ),
    },
    {
        "disclosure_id": "adjusted_loss_interpretation",
        "category": "outcome metric",
        "severity": "important",
        "title": "Adjusted loss may be negative",
        "statement": (
            "ATLAS-Tao unigram-adjusted loss is (NLL sum plus the unigram "
            "log-probability sum) divided by token count. The unigram term is "
            "negative, so valid finite adjusted losses may also be negative; "
            "lower remains better. It is not raw cross-entropy."
        ),
        "analysis_consequence": (
            "Fit and compare the declared adjusted metric at a common cohort. "
            "Use raw_nll for raw cross-entropy and do not exponentiate the "
            "adjusted value as perplexity."
        ),
    },
    {
        "disclosure_id": "analysis_ready_not_confirmatory_complete",
        "category": "analysis status",
        "severity": "important",
        "title": "Analysis-ready does not mean confirmatory-complete",
        "statement": (
            "The central tables are prepared for analysis, but the frozen C/D "
            "M1 score receipt and the subsequently declared A-D refit remain "
            "separate protocol stages."
        ),
        "analysis_consequence": (
            "Descriptive analysis may proceed. Do not claim confirmatory "
            "completion or interpret an A-D refit before the frozen C/D score "
            "is sealed."
        ),
    },
    {
        "disclosure_id": "actual_tokens",
        "category": "training design",
        "severity": "important",
        "title": "Actual token horizons are authoritative",
        "statement": (
            "Optimizer-step-aligned actual consumed tokens can differ slightly "
            "from nominal targets."
        ),
        "analysis_consequence": (
            "Use actual_tokens/actual_end_tokens for fits and plots; retain "
            "nominal targets only as configured design labels."
        ),
    },
    {
        "disclosure_id": "trunk_branch_reuse",
        "category": "compute accounting",
        "severity": "important",
        "title": "WSD trunk/branch reuse",
        "statement": (
            "A fixed-learning-rate trunk can be shared by multiple schedule-aware "
            "cooldown branches."
        ),
        "analysis_consequence": (
            "Use readout_compute_attribution only for logical per-readout cost. "
            "Use unique_physical_job_compute, unique_physical_study_compute, and "
            "physical_compute_categories for additive physical totals."
        ),
    },
    {
        "disclosure_id": "exact_language_exposure",
        "category": "data accounting",
        "severity": "important",
        "title": "Exact per-language exposure",
        "statement": (
            "Realized mixture exposure comes from exact blend replay and may not "
            "equal the legacy nominal D/K proxy."
        ),
        "analysis_consequence": (
            "Use data_exposure for language-level reuse and effective-epoch analyses."
        ),
    },
    {
        "disclosure_id": "measurement_dependence",
        "category": "statistical scope",
        "severity": "limitation",
        "title": "Nested repeated measurements",
        "statement": (
            "Languages are nested in readouts and checkpoint readouts are nested "
            "in trajectories."
        ),
        "analysis_consequence": (
            "Do not treat all language/checkpoint rows as IID observations."
        ),
    },
    {
        "disclosure_id": "gpu_board_energy",
        "category": "compute accounting",
        "severity": "limitation",
        "title": "GPU-board energy boundary",
        "statement": (
            "Integrated nvidia-smi power telemetry estimates GPU-board energy, "
            "not total system or facility energy."
        ),
        "analysis_consequence": (
            "Label it board energy and do not add PUE/carbon claims without "
            "separate system-boundary data."
        ),
    },
    {
        "disclosure_id": "recovery_candidates",
        "category": "validity",
        "severity": "important",
        "title": "Recovery copies are not replicates",
        "statement": (
            "Primary, recovery, and offload evaluation artifacts can represent "
            "the same checkpoint/readout."
        ),
        "analysis_consequence": (
            "Use the selected canonical candidate and retain all alternatives "
            "only for provenance and validity auditing."
        ),
    },
    {
        "disclosure_id": "provenance_classes",
        "category": "provenance",
        "severity": "important",
        "title": "Observed/configured/planned/derived are distinct",
        "statement": (
            "Runtime observations, frozen settings, forecasts, and recomputed "
            "values have different evidentiary status."
        ),
        "analysis_consequence": (
            "Preserve provenance_class and avoid presenting planned/configured "
            "values as observed results."
        ),
    },
    {
        "disclosure_id": "excel_is_derived",
        "category": "data access",
        "severity": "important",
        "title": "Parquet is canonical",
        "statement": (
            "The Excel workbook is a convenience view. High-cardinality tables "
            "may be intentionally omitted from it."
        ),
        "analysis_consequence": (
            "Use the Parquet tables for complete analysis and consult "
            "ExcelCoverage for every workbook omission."
        ),
    },
)


_EXCEL_SHEET_NAMES: dict[str, str] = {
    "release_summary": "ReleaseSummary",
    "studies": "Protocols",
    "source_roots": "SourceRoots",
    "model_configs": "ModelConfigs",
    "trajectories": "Trajectories",
    "readouts": "Readouts",
    "validation_aggregates": "ValidationAgg",
    "validation_language": "ValidationLanguage",
    "validation_candidates": "ValidationCandidates",
    "validation_scope": "ValidationScope",
    "training_jobs": "TrainingJobs",
    "training_configs": "TrainingConfigs",
    "training_segments": "TrainingSegments",
    "runtime_parameters": "RuntimeParameters",
    "training_summary": "TrainingSummary",
    "dataset_release": "DatasetRelease",
    "corpus_languages": "CorpusLanguages",
    "data_exposure": "DataExposure",
    "job_attempts": "JobAttempts",
    "slurm_allocations": "SlurmAllocations",
    "compute_summary": "ComputeSummary",
    "validity_events": "ValidityEvents",
    "qa_issues": "QAIssues",
    "analysis_registry": "AnalysisRegistry",
    "launch_provenance": "LaunchProvenance",
    "execution_inputs": "ExecutionInputs",
    "hardware_contract": "HardwareContract",
    "languages": "Languages",
    "language_tier_membership": "LanguageTiers",
    "language_metadata_sources": "LanguageMetaSources",
    "validation_cohorts": "ValidationCohorts",
    "readout_compute_attribution": "ReadoutCompute",
    "unique_physical_job_compute": "PhysicalJobCompute",
    "unique_physical_study_compute": "PhysicalStudyCompute",
    "physical_compute_categories": "ComputeCategories",
}

_ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def build_table_catalog(
    tables: Mapping[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Return the release table catalog, optionally enriched with observed schema.

    ``tables`` may contain pandas DataFrames.  Unknown supplied tables are
    retained in the returned catalog and marked ``uncatalogued`` so a new table
    cannot silently disappear from the data dictionary.
    """

    rows = [deepcopy(row) for row in _TABLE_CATALOG]
    by_name = {row["table_name"]: row for row in rows}
    if tables:
        for name, table in tables.items():
            if name not in by_name:
                row = _catalog_entry(
                    name,
                    "uncatalogued",
                    "not declared",
                    "Supplied release table without a registered table contract.",
                    (),
                    (),
                    (
                        "Add an explicit catalog entry before treating this table "
                        "as publication-ready."
                    ),
                    excel_default=False,
                )
                rows.append(row)
                by_name[name] = row
            row = by_name[name]
            columns = [str(column) for column in getattr(table, "columns", ())]
            dtypes = [str(dtype) for dtype in getattr(table, "dtypes", ())]
            row["row_count"] = int(len(table))
            row["rows"] = int(len(table))
            row["column_count"] = len(columns)
            row["observed_columns"] = columns
            row["observed_dtypes"] = dict(zip(columns, dtypes, strict=False))
            row["missing_key_columns"] = [
                column for column in row["key_columns"] if column not in columns
            ]
            primary_key = list(row["primary_key"])
            row["missing_primary_key_columns"] = [
                column for column in primary_key if column not in columns
            ]
            if not primary_key or row["missing_primary_key_columns"]:
                row["primary_key_null_row_count"] = None
                row["primary_key_duplicate_row_count"] = None
                row["primary_key_valid"] = False
            else:
                primary = table.loc[:, primary_key]
                row["primary_key_null_row_count"] = int(
                    primary.isna().any(axis=1).sum()
                )
                row["primary_key_duplicate_row_count"] = int(
                    primary.duplicated(keep=False).sum()
                )
                row["primary_key_valid"] = (
                    row["primary_key_null_row_count"] == 0
                    and row["primary_key_duplicate_row_count"] == 0
                )
            row["available"] = True
    supplied = set(tables or {})
    for row in rows:
        if row["table_name"] not in supplied:
            row.setdefault("row_count", None)
            row.setdefault("rows", None)
            row.setdefault("column_count", None)
            row.setdefault("observed_columns", [])
            row.setdefault("observed_dtypes", {})
            row.setdefault("missing_key_columns", [])
            row.setdefault("missing_primary_key_columns", [])
            row.setdefault("primary_key_null_row_count", None)
            row.setdefault("primary_key_duplicate_row_count", None)
            row.setdefault("primary_key_valid", None)
            row["available"] = False
    return rows


def build_data_dictionary(
    tables: Mapping[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Alias with publication-facing naming for :func:`build_table_catalog`."""

    return build_table_catalog(tables)


def _normalise_disclosures(
    disclosures: (Mapping[str, Any] | Sequence[Mapping[str, Any]] | None) = None,
) -> list[dict[str, str]]:
    merged = {row["disclosure_id"]: deepcopy(row) for row in _DEFAULT_DISCLOSURES}
    extra: list[Mapping[str, Any]]
    if disclosures is None:
        extra = []
    elif isinstance(disclosures, Mapping):
        extra = []
        for key, value in disclosures.items():
            if isinstance(value, Mapping):
                extra.append({"disclosure_id": key, **value})
            else:
                extra.append(
                    {
                        "disclosure_id": key,
                        "statement": str(value),
                    }
                )
    else:
        extra = list(disclosures)
    for index, supplied in enumerate(extra, start=1):
        disclosure_id = str(supplied.get("disclosure_id", f"additional_{index}"))
        base = merged.get(
            disclosure_id,
            {
                "disclosure_id": disclosure_id,
                "category": "additional",
                "severity": "important",
                "title": disclosure_id.replace("_", " ").title(),
                "statement": "",
                "analysis_consequence": "",
            },
        )
        base.update(
            {
                str(key): str(value)
                for key, value in supplied.items()
                if value is not None
            }
        )
        merged[disclosure_id] = base
    return list(merged.values())


def _metadata_value(
    metadata: Mapping[str, Any],
    *keys: str,
    default: Any = "not supplied",
) -> Any:
    for key in keys:
        if key in metadata and metadata[key] is not None:
            return metadata[key]
    return default


def _markdown_cell(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value).replace("|", r"\|").replace("\n", " ")


def _count_rows(
    table_counts: Mapping[str, Any] | None,
) -> list[tuple[str, Any]]:
    if not table_counts:
        return []
    rows: list[tuple[str, Any]] = []
    for name, value in sorted(table_counts.items()):
        if isinstance(value, Mapping):
            for child, count in sorted(value.items()):
                rows.append((f"{name}.{child}", count))
        else:
            rows.append((str(name), value))
    return rows


def build_neurips_checklist(
    release_metadata: Mapping[str, Any],
    table_counts: Mapping[str, Any] | None = None,
    disclosures: (Mapping[str, Any] | Sequence[Mapping[str, Any]] | None) = None,
    excel_omissions: Sequence[Mapping[str, Any]] | None = None,
) -> list[dict[str, str]]:
    """Build machine-readable NeurIPS-style release checklist evidence.

    The output deliberately distinguishes release support from the paper's
    eventual yes/no answer.  A data release cannot certify manuscript claims,
    citations, ethics statements, or accessibility on the authors' behalf.
    """

    counts = dict(table_counts or {})
    disclosure_rows = _normalise_disclosures(disclosures)
    disclosure_ids = {row["disclosure_id"] for row in disclosure_rows}
    physical_job_rows = int(counts.get("unique_physical_job_compute", 0) or 0)
    physical_study_rows = int(counts.get("unique_physical_study_compute", 0) or 0)
    physical_category_rows = int(counts.get("physical_compute_categories", 0) or 0)
    attribution_rows = int(counts.get("readout_compute_attribution", 0) or 0)
    launch_rows = int(counts.get("launch_provenance", 0) or 0)
    language_rows = int(counts.get("languages", 0) or 0)
    validation_language_rows = int(counts.get("validation_language", 0) or 0)
    cohort_rows = int(counts.get("validation_cohorts", 0) or 0)
    exposure_rows = int(counts.get("data_exposure", 0) or 0)
    trajectory_rows = int(
        counts.get(
            "trajectories",
            release_metadata.get("independent_trajectories", 0),
        )
        or 0
    )
    analysis_status = str(
        release_metadata.get(
            "confirmatory_analysis_status",
            "analysis_ready_not_confirmatory_complete",
        )
    )
    analysis_boundary = str(
        release_metadata.get("analysis_boundary")
        or (
            "Central analysis tables are available. The frozen C/D M1 score "
            "receipt and the subsequent declared A-D refit remain separate "
            "pending protocol stages."
        )
    )
    analysis_next_action = str(
        release_metadata.get("analysis_next_action")
        or (
            "Use the release for descriptive analysis without claiming "
            "confirmatory completion; seal the C/D score before the A-D refit."
        )
    )
    omissions = list(excel_omissions or [])
    physical_compute_tables_complete = bool(
        physical_job_rows
        and physical_study_rows == 1
        and physical_category_rows
        and attribution_rows
    )
    language_tables_complete = bool(
        language_rows == 200 and cohort_rows and exposure_rows
    )
    if excel_omissions is None:
        excel_coverage_evidence = (
            "The workbook writer emits one ExcelCoverage record per canonical "
            "table and never silently truncates a selected table."
        )
    else:
        excel_coverage_evidence = (
            f"{len(omissions)} Excel omissions are explicitly listed with their "
            "canonical Parquet paths; no selected table is silently truncated."
        )
    checklist = [
        {
            "checklist_id": "claims_limitations",
            "topic": "Claims and limitations",
            "release_support": "documented_with_limitations",
            "paper_answer": "supported; manuscript confirmation required",
            "evidence": (
                f"{len(disclosure_rows)} release disclosures, including "
                "single-seed and repeated-measure limitations."
            ),
            "action_required": (
                "Align each manuscript claim with a released outcome and carry "
                "the material limitations into the paper."
            ),
        },
        {
            "checklist_id": "analysis_status",
            "topic": "Analysis and confirmatory status",
            "release_support": analysis_status,
            "paper_answer": "confirmatory completion pending",
            "evidence": analysis_boundary,
            "action_required": analysis_next_action,
        },
        {
            "checklist_id": "assumptions_scope",
            "topic": "Assumptions and scope",
            "release_support": "documented",
            "paper_answer": "supported; manuscript confirmation required",
            "evidence": (
                "Table grains, provenance classes, actual-token policy, and the "
                "logical readout-attribution versus deduplicated physical-compute "
                "boundary are documented."
            ),
            "action_required": "State analysis estimands and independence assumptions.",
        },
        {
            "checklist_id": "theory",
            "topic": "Theoretical claims and proofs",
            "release_support": "not_applicable_to_empirical_release",
            "paper_answer": "N/A unless the manuscript adds theoretical claims",
            "evidence": "This artifact is an empirical data release.",
            "action_required": "Reassess at manuscript freeze.",
        },
        {
            "checklist_id": "reproducibility",
            "topic": "Experimental reproducibility",
            "release_support": "documented",
            "paper_answer": "supported; manuscript confirmation required",
            "evidence": (
                f"{trajectory_rows} independent trajectory rows and {launch_rows} "
                "execution-bound launch rows, plus resolved configurations, "
                "fixed-input hashes, source lineage, and analysis registry."
            ),
            "action_required": (
                "Publish accessible instructions and verify all referenced paths "
                "are portable or accompanied by retrieval guidance."
            ),
        },
        {
            "checklist_id": "data_code_access",
            "topic": "Data and code access",
            "release_support": "release_artifacts_prepared",
            "paper_answer": "manual availability confirmation required",
            "evidence": (
                "Canonical Parquet tables, manifest/hash inventory, data "
                "dictionary, and a derived workbook are specified."
            ),
            "action_required": (
                "Provide the anonymized or public archive URL, license, and "
                "artifact-access instructions in the submission."
            ),
        },
        {
            "checklist_id": "training_details",
            "topic": "Training and evaluation details",
            "release_support": "documented",
            "paper_answer": "supported",
            "evidence": (
                "Model configs, protocol seeds, job/segment topology, training "
                "metrics, readouts, and validation aggregates are cataloged."
            ),
            "action_required": "Cite the exact release version in the manuscript.",
        },
        {
            "checklist_id": "statistical_significance",
            "topic": "Statistical variation and uncertainty",
            "release_support": "documented_with_limitation",
            "paper_answer": "limitation disclosed",
            "evidence": (
                f"{trajectory_rows} independent full trajectories are available "
                f"with one configured seed per cell; the {validation_language_rows} language "
                "rows are nested repeated measurements."
            ),
            "action_required": (
                "Do not use language/checkpoint row counts as a substitute for "
                "independent seed replication; report the hierarchical design."
            ),
        },
        {
            "checklist_id": "compute",
            "topic": "Compute resources and efficiency",
            "release_support": (
                "documented"
                if physical_compute_tables_complete
                else "table_contract_documented_row_counts_unverified"
            ),
            "paper_answer": (
                "supported"
                if physical_compute_tables_complete
                else "row-count verification required"
            ),
            "evidence": (
                f"{physical_job_rows} unique physical-job rows, "
                f"{physical_study_rows} physical-study total row, "
                f"{physical_category_rows} workload-category rows, and "
                f"{attribution_rows} non-additive logical readout-attribution rows."
            ),
            "action_required": (
                "Report hardware, count, wall time/GPU-hours, and utilization "
                "boundary; never sum logical readout attribution as physical cost, "
                "and disclose that GPU-board energy is not facility energy."
            ),
        },
        {
            "checklist_id": "datasets",
            "topic": "Dataset documentation and provenance",
            "release_support": (
                "documented"
                if language_tables_complete
                else "table_contract_documented_row_counts_unverified"
            ),
            "paper_answer": (
                "supported"
                if language_tables_complete
                else "row-count verification required"
            ),
            "evidence": (
                f"{language_rows} canonical language rows, {cohort_rows} sealed "
                f"validation-cohort rows, and {exposure_rows} exact or explicitly "
                "censored language-exposure rows with source/corpus lineage."
            ),
            "action_required": (
                "Include corpus motivation, composition, licensing, filtering, "
                "deduplication, contamination, and access considerations."
            ),
        },
        {
            "checklist_id": "human_subjects",
            "topic": "Human subjects or crowdsourcing",
            "release_support": "not_assessed_by_release",
            "paper_answer": "manual ethics determination required",
            "evidence": (
                "Training/evaluation records do not by themselves establish the "
                "human-subject status of source corpora."
            ),
            "action_required": (
                "Authors must determine whether consent, compensation, IRB, or "
                "personally identifiable information requirements apply."
            ),
        },
        {
            "checklist_id": "broader_impacts",
            "topic": "Broader impacts, safeguards, and licenses",
            "release_support": "not_assessed_by_release",
            "paper_answer": "manual manuscript confirmation required",
            "evidence": "Technical release metadata cannot establish societal impact.",
            "action_required": (
                "Document intended use, limitations, multilingual risks, licenses, "
                "and foreseeable misuse in the paper."
            ),
        },
        {
            "checklist_id": "excel_coverage",
            "topic": "Convenience-view completeness",
            "release_support": "documented",
            "paper_answer": "N/A to paper; required for artifact usability",
            "evidence": excel_coverage_evidence,
            "action_required": (
                "Consult ExcelCoverage and its canonical_parquet_path before "
                "workbook-only analysis; use Parquet for complete analysis."
            ),
        },
    ]
    if "single_seed" not in disclosure_ids:
        raise ValueError("The single-seed disclosure may not be removed.")
    overrides = release_metadata.get("checklist_overrides", {})
    if isinstance(overrides, Mapping):
        for row in checklist:
            override = overrides.get(row["checklist_id"])
            if isinstance(override, Mapping):
                for key, value in override.items():
                    if key != "checklist_id":
                        row[str(key)] = str(value)
    return checklist


def render_release_readme(
    release_metadata: Mapping[str, Any],
    table_counts: Mapping[str, Any] | None = None,
    disclosures: (Mapping[str, Any] | Sequence[Mapping[str, Any]] | None) = None,
    *,
    excel_omissions: Sequence[Mapping[str, Any]] | None = None,
    catalog: Sequence[Mapping[str, Any]] | None = None,
    checklist: Sequence[Mapping[str, Any]] | None = None,
) -> str:
    """Render a publication-facing README for a sealed or candidate release."""

    disclosure_rows = _normalise_disclosures(disclosures)
    catalog_rows = list(catalog or build_table_catalog())
    checklist = list(
        checklist
        or build_neurips_checklist(
            release_metadata,
            table_counts,
            disclosure_rows,
            excel_omissions,
        )
    )
    release_name = _metadata_value(
        release_metadata, "release_name", "release_id", default="Stage 1 A-D"
    )
    schema = _metadata_value(release_metadata, "schema", "schema_version")
    created = _metadata_value(
        release_metadata,
        "built_at_utc",
        "created_at_utc",
        "generated_at_utc",
        "created_at",
    )
    status = _metadata_value(
        release_metadata, "status", "release_status", default="candidate"
    )
    analysis_status = _metadata_value(
        release_metadata,
        "analysis_status",
        "confirmatory_analysis_status",
        default="analysis_ready_not_confirmatory_complete",
    )
    analysis_boundary = str(
        release_metadata.get("analysis_boundary")
        or (
            "The frozen C/D M1 score and subsequent declared A-D refit remain "
            "separate protocol stages."
        )
    )
    analysis_next_action = str(
        release_metadata.get("analysis_next_action")
        or (
            "Use this release descriptively until the frozen C/D score is "
            "sealed, then perform the declared A-D refit."
        )
    )
    trajectory_count = int(
        (table_counts or {}).get("trajectories")
        or release_metadata.get("independent_trajectories")
        or 64
    )
    checksum = _metadata_value(
        release_metadata,
        "release_sha256",
        "manifest_sha256",
        "manifest_integrity_sha256",
        default="see release_manifest.json (manifest_integrity_sha256)",
    )
    count_lines = _count_rows(table_counts)
    if excel_omissions is None:
        readme_excel_omissions = [
            {
                "table_name": row["table_name"],
                "canonical_parquet_path": f"tables/{row['table_name']}.parquet",
                "status": "canonical_parquet_only_by_default",
                "reason": "not selected by the compact workbook default policy",
            }
            for row in catalog_rows
            if not row.get("excel_default", False)
        ]
        excel_coverage_basis = "catalog policy; verify final ExcelCoverage sheet"
    else:
        readme_excel_omissions = list(excel_omissions)
        excel_coverage_basis = "generated workbook coverage report"
    lines = [
        f"# {release_name}: central Stage 1 A--D analysis-ready data release",
        "",
        (
            "This versioned release centralizes the Stage 1 A-D multilingual "
            "pretraining ladder as analysis-ready data. It preserves execution-"
            "bound model/configuration detail, realized training and validation "
            "outcomes, exact data exposure, scheduler/compute accounting, "
            "telemetry, validity decisions, and artifact lineage. Analysis-ready "
            f"does not mean confirmatory-complete: {analysis_boundary}"
        ),
        "",
        "## Release identity",
        "",
        "| Field | Value |",
        "| --- | --- |",
        f"| Release | {_markdown_cell(release_name)} |",
        f"| Schema | {_markdown_cell(schema)} |",
        f"| Status | {_markdown_cell(status)} |",
        f"| Scientific analysis status | {_markdown_cell(analysis_status)} |",
        f"| Created | {_markdown_cell(created)} |",
        f"| Release/manifest SHA-256 | {_markdown_cell(checksum)} |",
        "",
        "## Canonical data and convenience views",
        "",
        (
            "Typed Parquet tables and the hash-sealed manifest are the canonical "
            "source of truth. The Excel workbook is a derived convenience view "
            "for inspection and compact analyses. Every canonical table has an "
            "`ExcelCoverage` record with its Parquet path. A table that violates "
            "an Excel limit is omitted in full and reported; it is never silently "
            "row- or column-truncated. Never infer that an omitted workbook sheet "
            "is absent from the release."
        ),
        "",
        "Minimal Python loading pattern:",
        "",
        "```python",
        "from pathlib import Path",
        "import pandas as pd",
        "",
        f'release = Path("result_analysis/{release_name}")',
        'readouts = pd.read_parquet(release / "tables/readouts.parquet")',
        (
            "validation = pd.read_parquet("
            'release / "tables/validation_language.parquet")'
        ),
        "```",
        "",
        "## Study design and unit of analysis",
        "",
        (
            f"The independent statistical unit is `trajectory_id`: {trajectory_count} "
            "full training trajectories in the A-D design. Each design cell has "
            "one configured seed. Schedule-aware checkpoints within a trajectory "
            "and languages within a checkpoint are nested repeated measurements. "
            "They must not be treated as independent seed replicates or used to "
            "make seed-general claims."
        ),
        "",
        (
            "Training uses explicit trunk/branch topology where a shared "
            "fixed-learning-rate trunk can feed multiple cooldown branches. "
            "`readout_compute_attribution` assigns a logical trunk prefix plus "
            "branch cost to each readout and is deliberately non-additive. Use "
            "`unique_physical_job_compute`, `unique_physical_study_compute`, and "
            "`physical_compute_categories` for additive physical totals; shared "
            "trunks and attempts are counted once. Actual optimizer-step-aligned "
            "consumed tokens, rather than nominal targets, are authoritative for "
            "training-horizon comparisons."
        ),
        "",
        "## Release contents",
        "",
    ]
    if count_lines:
        lines.extend(
            [
                "| Table/count | Rows or value |",
                "| --- | ---: |",
                *[
                    f"| `{_markdown_cell(name)}` | {_markdown_cell(value)} |"
                    for name, value in count_lines
                ],
            ]
        )
    else:
        lines.append(
            "Counts were not supplied to this README renderer; consult the manifest."
        )
    lines.extend(
        [
            "",
            "## Table catalog",
            "",
            "| Table | Grain | Purpose | Excel default |",
            "| --- | --- | --- | --- |",
        ]
    )
    for row in catalog_rows:
        if row.get("available") is False and catalog is not None:
            continue
        lines.append(
            "| `{}` | {} | {} | {} |".format(
                _markdown_cell(row["table_name"]),
                _markdown_cell(row["grain"]),
                _markdown_cell(row["purpose"]),
                "yes" if row.get("excel_default") else "Parquet only",
            )
        )
    lines.extend(
        [
            "",
            "The machine-readable data dictionary additionally records keys, "
            "observed columns/dtypes, row counts, and caveats.",
            "",
            "## Provenance and evidence classes",
            "",
            "| Class | Meaning |",
            "| --- | --- |",
        ]
    )
    for row in PROVENANCE_CLASSES:
        lines.append(
            f"| `{row['provenance_class']}` | {_markdown_cell(row['definition'])} |"
        )
    lines.extend(
        [
            "",
            "## Required disclosures and analysis consequences",
            "",
            "| Topic | Disclosure | Consequence |",
            "| --- | --- | --- |",
        ]
    )
    for row in disclosure_rows:
        lines.append(
            "| {} | {} | {} |".format(
                _markdown_cell(row["title"]),
                _markdown_cell(row["statement"]),
                _markdown_cell(row["analysis_consequence"]),
            )
        )
    lines.extend(
        [
            "",
            "## Compute and environmental accounting",
            "",
            (
                "`hardware_contract` freezes the MFU denominator. Scheduler "
                "allocations, application log throughput/MFU, and GPU telemetry "
                "answer different questions and remain separate. "
                "`physical_compute_categories` reports deduplicated workload "
                "scope; logical readout attribution is not a physical total and "
                "must never be summed across readouts. Allocated GPU-hours do not "
                "imply utilization. Integrated GPU power telemetry is GPU-board "
                "energy only and excludes hosts, network, storage, cooling, and PUE."
            ),
            "",
            "## Validity, exclusions, and analysis status",
            "",
            (
                "Technical-validity decisions must be traceable to evidence and "
                "must not exclude a result merely because its loss is poor. "
                "Primary/recovery/offload copies of the same checkpoint are "
                "candidate artifacts, not replicates. This release is analysis-"
                f"ready, not confirmatory-complete. {analysis_boundary} "
                f"Next action: {analysis_next_action} "
                "Observed/configured/planned and confirmatory/exploratory states "
                "must remain separate in `analysis_registry`; artifact build "
                "completeness is not scientific confirmatory completion."
            ),
            "",
            "## NeurIPS-style artifact checklist",
            "",
            "| Topic | Release support | Remaining action |",
            "| --- | --- | --- |",
        ]
    )
    for row in checklist:
        lines.append(
            "| {} | {} | {} |".format(
                _markdown_cell(row["topic"]),
                _markdown_cell(row["release_support"]),
                _markdown_cell(row["action_required"]),
            )
        )
    lines.extend(
        [
            "",
            (
                "This table is evidence support for the authors. It does not "
                "replace the paper-level checklist or ethics/licensing review."
            ),
            "",
            "## Excel coverage",
            "",
        ]
    )
    omissions = readme_excel_omissions
    lines.append(f"Coverage basis: {excel_coverage_basis}.")
    if omissions:
        lines.extend(
            [
                "| Canonical table | Canonical Parquet path | Workbook status | Reason |",
                "| --- | --- | --- | --- |",
            ]
        )
        for row in omissions:
            lines.append(
                "| `{}` | `{}` | {} | {} |".format(
                    _markdown_cell(row.get("table_name", "unknown")),
                    _markdown_cell(
                        row.get(
                            "canonical_parquet_path",
                            f"tables/{row.get('table_name', 'unknown')}.parquet",
                        )
                    ),
                    _markdown_cell(row.get("status", "omitted")),
                    _markdown_cell(row.get("reason", "not supplied")),
                )
            )
    else:
        lines.append(
            "No tables are omitted in the supplied coverage report. Consult "
            "`ExcelCoverage` to verify every canonical table was included complete."
        )
    lines.extend(
        [
            "",
            "## Path portability and generation code",
            "",
            (
                "Canonical release paths are relative to this release directory. "
                "Absolute HPC paths retained in provenance fields identify the "
                "build-host source artifacts; they are evidence locators, not "
                "portable runtime dependencies. Remap them with `source_roots` "
                "and verify the recorded hashes. The exact release-builder, "
                "extractor, notebook-builder, and focused test sources used for "
                "this build are copied under `code/` and sealed in the manifest's "
                "generation-code inventory."
            ),
            "",
            "## Integrity and citation",
            "",
            (
                "Verify the release manifest and file SHA-256 hashes before "
                "analysis. Cite the exact release name/schema and checksum; do "
                "not cite a mutable working directory as though it were sealed."
            ),
            "",
        ]
    )
    return "\n".join(lines)


def render_readme(
    release_metadata: Mapping[str, Any],
    catalog: Sequence[Mapping[str, Any]] | None = None,
    checklist: Sequence[Mapping[str, Any]] | None = None,
    *,
    table_counts: Mapping[str, Any] | None = None,
    disclosures: (Mapping[str, Any] | Sequence[Mapping[str, Any]] | None) = None,
    excel_omissions: Sequence[Mapping[str, Any]] | None = None,
) -> str:
    """Render the README using optionally precomputed presentation records."""

    return render_release_readme(
        release_metadata,
        table_counts,
        disclosures,
        excel_omissions=excel_omissions,
        catalog=catalog,
        checklist=checklist,
    )


def _json_excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        )
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _prepare_excel_frame(frame: Any) -> tuple[Any, dict[str, int]]:
    prepared = frame.copy()
    transformations = {
        "timezone_datetime_to_iso": 0,
        "nested_value_to_json": 0,
        "empty_string_to_blank": 0,
        "illegal_character_replacements": 0,
        "formula_escaped_cells": 0,
    }
    for column in prepared.columns:
        series = prepared[column]
        dtype = getattr(series, "dtype", None)
        if getattr(dtype, "tz", None) is not None:
            transformations["timezone_datetime_to_iso"] += int(series.notna().sum())
            prepared[column] = series.map(
                lambda value: value.isoformat()
                if value is not None and not _is_missing_scalar(value)
                else value
            )
            series = prepared[column]
        if str(dtype).startswith(("period", "interval")):
            prepared[column] = series.astype("string")
            series = prepared[column]
        if str(getattr(series, "dtype", "")) in {"object", "string", "str"}:

            def convert(value: Any) -> Any:
                if value is None or _is_missing_scalar(value):
                    return value
                if isinstance(value, (dict, list, tuple, set)):
                    transformations["nested_value_to_json"] += 1
                value = _json_excel_value(value)
                if not isinstance(value, str):
                    return value
                if value == "":
                    transformations["empty_string_to_blank"] += 1
                    return None
                cleaned, replacements = _ILLEGAL_EXCEL_CHARACTERS.subn(
                    "\N{REPLACEMENT CHARACTER}", value
                )
                transformations["illegal_character_replacements"] += replacements
                if cleaned.startswith(("=", "+", "-", "@")):
                    transformations["formula_escaped_cells"] += 1
                    return "'" + cleaned
                return cleaned

            prepared[column] = series.map(convert)
    return prepared, transformations


def _is_missing_scalar(value: Any) -> bool:
    try:
        result = value != value
    except (TypeError, ValueError):
        return False
    return isinstance(result, bool) and result


def _overlong_cell_count(frame: Any) -> int:
    count = 0
    for column in frame.columns:
        series = frame[column]
        if str(getattr(series, "dtype", "")) not in {"object", "string", "str"}:
            continue
        count += int(
            series.map(
                lambda value: isinstance(value, str)
                and len(value) > EXCEL_MAX_CELL_CHARACTERS
            ).sum()
        )
    return count


def _flatten_metadata(
    metadata: Mapping[str, Any],
    prefix: str = "",
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for key, value in metadata.items():
        dotted = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, Mapping):
            rows.extend(_flatten_metadata(value, dotted))
        else:
            rendered = _json_excel_value(value)
            rows.append({"field": dotted, "value": str(rendered)})
    return rows


def _safe_sheet_name(table_name: str, used: set[str]) -> str:
    proposed = _EXCEL_SHEET_NAMES.get(table_name)
    if proposed is None:
        proposed = re.sub(r"[\[\]:*?/\\]", "_", table_name)[:31] or "Table"
    base = proposed[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        ending = f"_{suffix}"
        candidate = base[: 31 - len(ending)] + ending
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _stringify_dictionary_rows(
    rows: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for row in rows:
        result.append(
            {
                key: (
                    json.dumps(
                        value,
                        ensure_ascii=False,
                        sort_keys=True,
                        separators=(",", ":"),
                    )
                    if isinstance(value, (dict, list, tuple))
                    else value
                )
                for key, value in row.items()
            }
        )
    return result


def _style_workbook(workbook: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    wrap_headers = {
        "README",
        "Provenance",
        "Limitations",
        "DataDictionary",
        "NeurIPS",
        "ExcelCoverage",
    }
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.row_dimensions[1].height = 28
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        sample_rows = list(
            worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row, 250),
                values_only=True,
            )
        )
        for index in range(1, worksheet.max_column + 1):
            width = 10
            for row in sample_rows:
                value = row[index - 1] if index - 1 < len(row) else None
                if value is not None:
                    width = max(width, min(len(str(value)) + 2, 58))
            worksheet.column_dimensions[get_column_letter(index)].width = width
        if worksheet.title in wrap_headers:
            for row in worksheet.iter_rows(
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            worksheet.sheet_properties.tabColor = "5B9BD5"


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_excel_workbook(
    output_path: str | Path,
    tables: Mapping[str, Any],
    release_metadata: Mapping[str, Any],
    catalog: Sequence[Mapping[str, Any]] | None = None,
    checklist: Sequence[Mapping[str, Any]] | None = None,
    disclosures: (Mapping[str, Any] | Sequence[Mapping[str, Any]] | None) = None,
    *,
    include_tables: Sequence[str] | None = None,
    max_data_rows_per_sheet: int = EXCEL_MAX_DATA_ROWS,
) -> dict[str, Any]:
    """Write a styled, atomic, explicitly scoped Excel convenience workbook.

    No table is row-truncated. A selected table that exceeds an Excel limit is
    omitted in full and recorded on ``ExcelCoverage``. Every supplied table not
    selected for the workbook is also recorded there as canonical-Parquet-only.
    The returned report contains the same coverage records and workbook hash.
    """

    try:
        import pandas as pd
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise RuntimeError("write_excel_workbook requires pandas and openpyxl") from exc
    if not 1 <= max_data_rows_per_sheet <= EXCEL_MAX_DATA_ROWS:
        raise ValueError(
            f"max_data_rows_per_sheet must be between 1 and {EXCEL_MAX_DATA_ROWS}"
        )
    for name, frame in tables.items():
        if not isinstance(frame, pd.DataFrame):
            raise TypeError(f"tables[{name!r}] is not a pandas DataFrame")
    unknown_requested = set(include_tables or ()) - set(tables)
    if unknown_requested:
        raise KeyError(
            "Requested Excel tables were not supplied: "
            + ", ".join(sorted(unknown_requested))
        )

    catalog_rows = (
        [deepcopy(dict(row)) for row in catalog]
        if catalog is not None
        else build_table_catalog(tables)
    )
    catalog_by_name = {row["table_name"]: row for row in catalog_rows}
    missing_catalog = set(tables) - set(catalog_by_name)
    if missing_catalog:
        observed_catalog = {
            row["table_name"]: row for row in build_table_catalog(tables)
        }
        for name in sorted(missing_catalog):
            row = observed_catalog[name]
            catalog_rows.append(row)
            catalog_by_name[name] = row
    selected = (
        set(include_tables)
        if include_tables is not None
        else {
            name for name in tables if catalog_by_name[name].get("excel_default", False)
        }
    )
    catalog_order = {row["table_name"]: index for index, row in enumerate(catalog_rows)}
    coverage: list[dict[str, Any]] = []
    prepared_tables: dict[str, Any] = {}
    for table_name in sorted(
        tables, key=lambda name: (catalog_order.get(name, 10_000), name)
    ):
        frame = tables[table_name]
        record: dict[str, Any] = {
            "table_name": table_name,
            "canonical_format": "Parquet",
            "canonical_parquet_path": f"tables/{table_name}.parquet",
            "canonical_source_of_truth": True,
            "canonical_rows": int(len(frame)),
            "canonical_columns": int(len(frame.columns)),
            "requested_for_excel": table_name in selected,
            "excel_rows": 0,
            "excel_columns": 0,
            "complete_row_and_column_coverage": False,
            "omitted_from_excel": True,
            "status": "canonical_parquet_only",
            "reason": "not selected for the compact convenience workbook",
            "transformations": "",
        }
        if table_name not in selected:
            coverage.append(record)
            continue
        if not frame.columns.is_unique:
            record["status"] = "omitted_schema_limit"
            record["reason"] = "duplicate column names are unsafe in Excel"
            coverage.append(record)
            continue
        if len(frame.columns) > EXCEL_MAX_COLUMNS:
            record["status"] = "omitted_excel_column_limit"
            record["reason"] = (
                f"{len(frame.columns)} columns exceeds Excel's "
                f"{EXCEL_MAX_COLUMNS}-column limit"
            )
            coverage.append(record)
            continue
        if len(frame) > max_data_rows_per_sheet:
            record["status"] = "omitted_excel_row_limit"
            record["reason"] = (
                f"{len(frame)} rows exceeds configured no-truncation limit "
                f"{max_data_rows_per_sheet}"
            )
            coverage.append(record)
            continue
        prepared, transformations = _prepare_excel_frame(frame)
        overlong = _overlong_cell_count(prepared)
        if overlong:
            record["status"] = "omitted_excel_cell_limit"
            record["reason"] = (
                f"{overlong} cells exceed Excel's "
                f"{EXCEL_MAX_CELL_CHARACTERS}-character cell limit"
            )
            record["transformations"] = json.dumps(transformations, sort_keys=True)
            coverage.append(record)
            continue
        prepared_tables[table_name] = prepared
        record["status"] = "included_complete"
        record["reason"] = "all canonical rows and columns included"
        record["excel_rows"] = int(len(prepared))
        record["excel_columns"] = int(len(prepared.columns))
        record["complete_row_and_column_coverage"] = True
        record["omitted_from_excel"] = False
        record["transformations"] = json.dumps(transformations, sort_keys=True)
        coverage.append(record)

    covered_tables = {str(row["table_name"]) for row in coverage}
    if len(coverage) != len(tables) or covered_tables != set(tables):
        raise AssertionError(
            "ExcelCoverage must contain exactly one record per canonical table"
        )

    omissions = [row for row in coverage if row["status"] != "included_complete"]
    counts = {name: int(len(frame)) for name, frame in tables.items()}
    disclosure_rows = _normalise_disclosures(disclosures)
    analysis_boundary = str(
        release_metadata.get("analysis_boundary")
        or (
            "The frozen C/D M1 score receipt must precede the subsequently "
            "declared A-D refit."
        )
    )
    checklist_rows = list(
        checklist
        or build_neurips_checklist(
            release_metadata,
            counts,
            disclosure_rows,
            omissions,
        )
    )
    dictionary_rows = catalog_rows
    completeness_rows = [
        {
            "table_name": name,
            "rows": int(len(frame)),
            "columns": int(len(frame.columns)),
            "catalog_status": (
                "declared"
                if catalog_by_name[name]["category"] != "uncatalogued"
                else "uncatalogued"
            ),
            "missing_key_columns": ", ".join(
                catalog_by_name[name]["missing_key_columns"]
            ),
        }
        for name, frame in sorted(tables.items())
    ]
    provenance_rows = [
        *[dict(row) for row in PROVENANCE_CLASSES],
        {
            "provenance_class": "release_manifest",
            "definition": (
                "Release identity: "
                + str(
                    _metadata_value(
                        release_metadata,
                        "release_name",
                        "release_id",
                    )
                )
                + "; source snapshot: "
                + str(
                    _metadata_value(
                        release_metadata,
                        "source_snapshot_at_utc",
                        "accounting_snapshot_at_utc",
                        default="recorded in manifest",
                    )
                )
            ),
        },
    ]
    readme_rows = [
        {
            "section": "NOTICE",
            "item": "canonical format",
            "value": (
                "Parquet tables and the hash-sealed manifest are the canonical "
                "source of truth. This workbook is a derived convenience view."
            ),
        },
        {
            "section": "NOTICE",
            "item": "coverage",
            "value": (
                "Consult ExcelCoverage before workbook-only analysis. No selected "
                "table is silently row- or column-truncated; omissions are whole-"
                "table records with canonical Parquet paths."
            ),
        },
        {
            "section": "NOTICE",
            "item": "statistical unit",
            "value": (
                "The independent unit is trajectory_id: 64 full trajectories, "
                "with one configured seed per design cell. Language/checkpoint "
                "rows are nested repeated measurements."
            ),
        },
        {
            "section": "NOTICE",
            "item": "analysis status",
            "value": "Analysis-ready, not confirmatory-complete. " + analysis_boundary,
        },
        {
            "section": "NOTICE",
            "item": "compute aggregation",
            "value": (
                "ReadoutCompute is logical and non-additive. Use "
                "PhysicalJobCompute, PhysicalStudyCompute, and ComputeCategories "
                "for additive physical compute."
            ),
        },
        *[
            {
                "section": "RELEASE METADATA",
                "item": row["field"],
                "value": row["value"],
            }
            for row in _flatten_metadata(release_metadata)
        ],
    ]
    front_frames = {
        "README": pd.DataFrame(readme_rows),
        "Provenance": pd.DataFrame(provenance_rows),
        "Limitations": pd.DataFrame(disclosure_rows),
        "Completeness": pd.DataFrame(completeness_rows),
        "DataDictionary": pd.DataFrame(_stringify_dictionary_rows(dictionary_rows)),
        "NeurIPS": pd.DataFrame(checklist_rows),
        "ExcelCoverage": pd.DataFrame(coverage),
    }

    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + ".tmp.xlsx")
    if temporary.exists():
        temporary.unlink()
    used_names = {name.casefold() for name in front_frames}
    written_sheets: list[str] = []
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for sheet_name, frame in front_frames.items():
                frame.to_excel(writer, sheet_name=sheet_name, index=False)
                written_sheets.append(sheet_name)
            for table_name, frame in prepared_tables.items():
                sheet_name = _safe_sheet_name(table_name, used_names)
                frame.to_excel(writer, sheet_name=sheet_name, index=False)
                written_sheets.append(sheet_name)
            _style_workbook(writer.book)
            writer.book.properties.title = (
                str(
                    _metadata_value(
                        release_metadata,
                        "release_name",
                        "release_id",
                        default="Stage 1 A-D analytical release",
                    )
                )
                + " convenience workbook"
            )
            writer.book.properties.subject = (
                "Derived Excel view; canonical data are Parquet"
            )
            writer.book.properties.creator = "multilingual-megatron-study"
        temporary.replace(destination)
    except Exception:
        if temporary.exists():
            temporary.unlink()
        raise
    return {
        "path": str(destination),
        "sha256": _sha256_file(destination),
        "size_bytes": destination.stat().st_size,
        "sheets": written_sheets,
        "coverage": coverage,
        "omissions": omissions,
        "canonical_format": "Parquet",
        "canonical_source_of_truth": "Parquet tables plus hash-sealed manifest",
        "canonical_table_count": len(tables),
        "excel_coverage_record_count": len(coverage),
        "included_complete_table_count": len(prepared_tables),
        "omitted_table_count": len(omissions),
        "no_silent_truncation": True,
        "no_silent_row_truncation": True,
        "no_silent_column_truncation": True,
    }


__all__ = [
    "EXCEL_MAX_CELL_CHARACTERS",
    "EXCEL_MAX_COLUMNS",
    "EXCEL_MAX_DATA_ROWS",
    "PROVENANCE_CLASSES",
    "build_data_dictionary",
    "build_neurips_checklist",
    "build_table_catalog",
    "render_readme",
    "render_release_readme",
    "write_excel_workbook",
]
